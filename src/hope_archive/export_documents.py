"""Shared diary presentation with bounded, offline document renderers."""
from datetime import datetime
from enum import Enum
from html import escape
from io import BytesIO
import hashlib
import os
import re
from pathlib import Path
import zipfile

from PIL import Image, ImageOps
from .document_style import STYLES, BODY, paragraphs, CommentText, comment_text, LATIN_FONT, CJK_FONT
from .export_markdown import (diary_path, display_value, EMOTION_LABELS, WEATHER_LABELS,
                              resolve_media, person_name, export_diaries)


class ExportFormat(str, Enum):
    MARKDOWN = 'markdown'
    TEX = 'tex'
    PDF = 'pdf'
    DOCX = 'docx'
    TXT = 'txt'


def fit_image(width, height, max_width=450, max_height=600):
    if width <= 0 or height <= 0:
        raise ValueError('Invalid image dimensions')
    scale = min(max_width / width, max_height / height, 1)
    return width * scale, height * scale


def image_parts(path):
    """One oriented image, never slice a screenshot into oversized page fragments."""
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        if image.mode not in ('RGB', 'RGBA'):
            image = image.convert('RGB')
        data = BytesIO()
        image.save(data, format='PNG')
        yield data.getvalue(), image.size


def blocks(diary, manifest, archive):
    from .document import diary_blocks
    for kind, value in diary_blocks(diary):
        if kind in ('image', 'media'):
            media_kind, url = value
            local = resolve_media(url, manifest, archive)
            if kind == 'image' and local:
                try:
                    yield from (('image', part) for part in image_parts(local))
                except (OSError, ValueError, SyntaxError, Image.DecompressionBombError):
                    yield 'missing', '[本地媒体不可用]'
            else:
                yield 'missing' if not local else 'text', '[本地媒体不可用]' if not local else f'[{media_kind}：{local.name}]'
        elif kind == 'comment':
            yield 'comment', comment_text(*value)
        else:
            yield kind, value


def tex_escape(text):
    mapping = {'\\': r'\textbackslash{}', '{': r'\{', '}': r'\}', '$': r'\$',
               '&': r'\&', '#': r'\#', '%': r'\%', '_': r'\_',
               '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(mapping.get(char, char) for char in text)


def write_exclusive(path, data):
    if path.exists():
        if path.read_bytes() != data:
            raise FileExistsError('Existing export differs; refusing overwrite')
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('xb') as handle:
        handle.write(data)
    return True


def render_tex(items, path):
    parts = [r'\documentclass[UTF8,fontset=fandol]{ctexart}', r'\usepackage[a4paper,margin=25mm]{geometry}',
             r'\usepackage{graphicx}', r'\setlength{\parindent}{0pt}', r'\setlength{\parskip}{9pt}', r'\begin{document}']
    pending=[]
    def flush():
        for row in pdf_image_rows(pending, available_width=453.54):
            pictures=[]
            for data,(width,height) in row:
                name=hashlib.sha256(data).hexdigest()+'.png'
                write_exclusive(path.parent/'assets'/name,data)
                pictures.append(r'\includegraphics[width='+f'{width:.3f}pt,height={height:.3f}pt'+r',keepaspectratio]{assets/'+name+'}')
            parts.append(r'\par\noindent\makebox[\linewidth][c]{'+r'\hspace{.02\linewidth}'.join(pictures)+r'}\par\vspace{12pt}')
        pending.clear()
    for kind, value in items:
        if kind == 'image':
            pending.append(value)
            continue
        flush()
        if kind == 'page':
            parts.append(r'\newpage')
        elif kind == 'separator':
            parts.append(r'\par\bigskip\noindent\rule{\linewidth}{0.4pt}\par\bigskip')
        elif kind in ('heading', 'title', 'subheading'):
            parts.append((r'\section*{' if kind=='heading' else r'\subsection*{') + tex_escape(value) + '}')
        else:
            parts.append('\n\n'.join(tex_escape(line) + r'\par' for line in str(value).splitlines()))
    flush()
    return ('\n\n'.join(parts) + '\n\\end{document}\n').encode('utf-8')


def render_docx(items):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(56)
    for name in ('Normal', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        fonts = style.element.get_or_add_rPr().rFonts
        fonts.set(qn('w:eastAsia'), CJK_FONT)
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
    doc.core_properties.author = 'Hope Archive'
    doc.core_properties.last_modified_by = 'Hope Archive'
    doc.core_properties.created = doc.core_properties.modified = datetime(2000, 1, 1)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import RGBColor
    for name, semantic in [('Normal','text'),('Heading 1','heading'),('Heading 2','title')]:
        spec = STYLES[semantic]
        style = doc.styles[name]
        style.font.size = Pt(spec.size)
        style.font.bold = spec.bold
        fmt = style.paragraph_format
        fmt.line_spacing = Pt(spec.leading)
        fmt.space_before, fmt.space_after = Pt(spec.before), Pt(spec.after)
        fmt.keep_with_next = spec.keep_next
        fmt.widow_control = True
        doc.styles[name].font.color.rgb=RGBColor.from_string('222222')
        if spec.rule:
            border = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            for key, val in [('val','single'),('sz','4'),('color','CCCCCC'),('space','5')]: bottom.set(qn('w:' + key), val)
            border.append(bottom); style.element.get_or_add_pPr().append(border)
    usable = (section.page_width - section.left_margin - section.right_margin) / 12700
    pending=[]
    def add_image(paragraph,data,size):
        paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after=Pt(12)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.add_run().add_picture(BytesIO(data),width=Pt(size[0]),height=Pt(size[1]))
    def flush():
        for row in pdf_image_rows(pending, available_width=usable):
            if len(row)==1:
                add_image(doc.add_paragraph(),*row[0])
            else:
                from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
                from .image_layout import cell_fraction
                count = len(row)
                table=doc.add_table(rows=1,cols=count * 2 - 1)
                table.autofit=False
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                borders = OxmlElement('w:tblBorders')
                for edge in ('top','left','bottom','right','insideH','insideV'):
                    item = OxmlElement('w:' + edge); item.set(qn('w:val'), 'nil'); borders.append(item)
                table._tbl.tblPr.append(borders)
                widths=[usable * (cell_fraction(count, None) if i % 2 == 0 else .02) for i in range(count * 2 - 1)]
                for column,width in zip(table.columns,widths): column.width=Pt(width)
                for cell,width in zip(table.rows[0].cells,widths):
                    cell.width=Pt(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'0');edge.set(qn('w:type'),'dxa');margins.append(edge)
                    cell._tc.get_or_add_tcPr().append(margins)
                props=table.rows[0]._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
                for index,(data,size) in enumerate(row): add_image(table.cell(0,index * 2).paragraphs[0],data,size)
        pending.clear()
    for kind, value in items:
        if kind=='image':
            pending.append(value)
            continue
        flush()
        if kind == 'page':
            doc.add_page_break()
        elif kind == 'separator':
            paragraph = doc.add_paragraph()
            border = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            for key, val in [('val','single'),('sz','4'),('color','CCCCCC')]: bottom.set(qn('w:' + key), val)
            border.append(bottom); paragraph._p.get_or_add_pPr().append(border)
        elif kind in ('heading', 'title', 'subheading'):
            doc.add_heading(value, level=1 if kind == 'heading' else 2)
        else:
            chunks = paragraphs(value.body if isinstance(value, CommentText) else value)
            if isinstance(value, CommentText) and not chunks: chunks = ['']
            for index, chunk in enumerate(chunks):
                paragraph = doc.add_paragraph()
                if index == 0 and isinstance(value, CommentText):
                    for label, bold in value.runs: paragraph.add_run(label).bold = bold
                paragraph.add_run(chunk)
    flush()
    result = BytesIO()
    doc.save(result)
    # Stable timestamps make repeated identical exports safely skippable.
    stable = BytesIO()
    with zipfile.ZipFile(result) as source, zipfile.ZipFile(stable, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    return stable.getvalue()


from .image_layout import PDF_PAGE, PDF_MARGIN, PDF_BODY_WIDTH, PDF_IMAGE_GAP, image_rows as pdf_image_rows


def render_pdf(items):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Image as PDFImage, Spacer, PageBreak, Table, TableStyle, HRFlowable
    font = 'HopeChinese'
    latin = None
    font_dir = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    symbol = None
    if (font_dir / 'seguisym.ttf').is_file():
        symbol = 'HopeSymbols'
        if symbol not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(symbol, str(font_dir / 'seguisym.ttf')))
    if (font_dir / 'segoeui.ttf').is_file() and (font_dir / 'segoeuib.ttf').is_file():
        latin = 'HopeLatin'
        if latin not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(latin, str(font_dir / 'segoeui.ttf')))
            pdfmetrics.registerFont(TTFont(latin+'Bold', str(font_dir / 'segoeuib.ttf')))
    if font not in pdfmetrics.getRegisteredFontNames():
        font_dir = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
        system_font = font_dir / 'msyh.ttc'
        if not system_font.is_file(): system_font = font_dir / 'simsun.ttc'
        if system_font.is_file():
            pdfmetrics.registerFont(TTFont(font, str(system_font), subfontIndex=0))
            bold_file = font_dir / 'msyhbd.ttc'
            bold = font + 'Bold'
            pdfmetrics.registerFont(TTFont(bold, str(bold_file if bold_file.is_file() else system_font), subfontIndex=0))
            pdfmetrics.registerFontFamily(font, normal=font, bold=bold, italic=font, boldItalic=bold)
        else:
            font = 'STSong-Light'
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(UnicodeCIDFont(font))
    regular = ParagraphStyle('Diary', fontName=font, fontSize=BODY.size, leading=BODY.leading, spaceAfter=BODY.after, wordWrap='CJK', textColor=HexColor('#222222'))
    styles = {kind: ParagraphStyle(kind, parent=regular, fontSize=spec.size, leading=spec.leading,
        spaceBefore=spec.before, spaceAfter=spec.after, keepWithNext=spec.keep_next,
        fontName=font + 'Bold' if spec.bold and font + 'Bold' in pdfmetrics.getRegisteredFontNames() else font)
        for kind, spec in STYLES.items()}
    def literal(text, bold=False):
        # Font selection is explicit; escape every user character before markup.
        result = []
        for run in re.findall(r'[\x20-\x7e]+|[^\x20-\x7e]+', str(text)):
            escaped = escape(run).replace('\n', '<br/>')
            if latin and all(ord(c) < 128 for c in run):
                escaped = '<font name="' + latin + ('Bold' if bold else '') + '">' + escaped + '</font>'
            elif bold:
                escaped = '<b>' + escaped + '</b>'
            if symbol:
                for glyph in ('↳', '☀'):
                    escaped = escaped.replace(glyph, '<font name="' + symbol + '">' + glyph + '</font>')
                escaped = escaped.replace('\ufe0f', '')
            result.append(escaped)
        return ''.join(result)

    class HeadingParagraph(Paragraph):
        def draw(self):
            super().draw()
            self.canv.saveState()
            self.canv.setStrokeColor(HexColor('#CCCCCC'))
            self.canv.setLineWidth(.5)
            self.canv.line(0, -5, self.width, -5)
            self.canv.restoreState()
    story, pending = [], []
    def flush_images():
        for row in pdf_image_rows(pending):
            images = [PDFImage(BytesIO(data), width=size[0], height=size[1], hAlign='CENTER') for data, size in row]
            if len(images) == 1:
                story.append(images[0])
            else:
                from .image_layout import cell_fraction
                cells, widths = [], []
                for index, item in enumerate(images):
                    if index: cells.append(''); widths.append(PDF_IMAGE_GAP)
                    cells.append(item); widths.append(PDF_BODY_WIDTH * cell_fraction(len(images), None))
                table = Table([cells], colWidths=widths, hAlign='CENTER')
                table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'), ('ALIGN',(0,0),(-1,-1),'CENTER'),
                    ('LEFTPADDING',(0,0),(-1,-1),0), ('RIGHTPADDING',(0,0),(-1,-1),0),
                    ('TOPPADDING',(0,0),(-1,-1),0), ('BOTTOMPADDING',(0,0),(-1,-1),0)]))
                story.append(table)
            story.append(Spacer(1, PDF_IMAGE_GAP))
        pending.clear()
    for kind, value in items:
        if kind == 'image':
            pending.append(value)
            continue
        flush_images()
        if kind == 'page':
            story.append(PageBreak())
        elif kind == 'separator':
            story.append(HRFlowable(width='100%', thickness=.5, color=HexColor('#CCCCCC'), spaceBefore=12, spaceAfter=12))
        else:
            # Paragraphs remain literal diary text; never interpret user HTML.
            text = str(value)
            chunks = paragraphs(value.body if isinstance(value, CommentText) else text)
            if isinstance(value, CommentText) and not chunks: chunks = ['']
            for index, chunk in enumerate(chunks):
                spec = STYLES.get(kind, BODY)
                markup = literal(chunk, spec.bold)
                if index == 0 and isinstance(value, CommentText):
                    markup = ''.join(literal(label, bold) for label, bold in value.runs) + markup
                paragraph_class = HeadingParagraph if spec.rule else Paragraph
                story.append(paragraph_class(markup, styles.get(kind, regular)))
    flush_images()
    target = BytesIO()
    document = BaseDocTemplate(target, pagesize=PDF_PAGE, rightMargin=PDF_MARGIN, leftMargin=PDF_MARGIN,
                      topMargin=54, bottomMargin=54, invariant=1, title='Hope Archive', author='Hope Archive')
    document.addPageTemplates(PageTemplate(id='diary', frames=[Frame(PDF_MARGIN,54,PDF_BODY_WIDTH,PDF_PAGE[1]-108,
        leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)]))
    document.build(story)
    return target.getvalue()


def render_items(items, path, format):
    """Shared output renderers for diaries, opened capsules and editable OCR pages."""
    format = ExportFormat(format)
    if format == ExportFormat.TEX: return render_tex(items, path)
    if format == ExportFormat.PDF: return render_pdf(items)
    if format == ExportFormat.DOCX: return render_docx(items)
    return '\n\n'.join(('\f' if format == ExportFormat.TXT else '---') if kind == 'page'
        else ('# ' if kind == 'heading' and format == ExportFormat.MARKDOWN else '') + str(value)
        for kind, value in items if kind != 'image').encode('utf-8')


def export_pages(pages, output, format, title='识图文字'):
    import uuid
    if not isinstance(pages, list) or not 1 <= len(pages) <= 200:
        raise ValueError('没有可导出的页面。')
    items = []
    for index, page in enumerate(pages):
        if not isinstance(page, dict) or not isinstance(page.get('text'), str) or len(page['text']) > 500_000:
            raise ValueError('页面内容无效。')
        if index: items.append(('page', ''))
        items.extend([('heading', f'{title} · {index + 1}'), ('text', page['text'])])
    format = ExportFormat(format)
    suffix = 'md' if format == ExportFormat.MARKDOWN else format.value
    path = Path(output) / f'{title}_{datetime.now():%Y%m%d-%H%M%S}_{uuid.uuid4().hex[:10]}.{suffix}'
    path.parent.mkdir(parents=True, exist_ok=True)
    write_exclusive(path, render_items(items, path, format))
    return {'path': str(path), 'pages': len(pages)}


def export_document(document, manifest, archive, output_dir, format='markdown'):
    format = ExportFormat(format)
    archive, output_dir = Path(archive).resolve(), Path(output_dir).resolve()
    if format == ExportFormat.MARKDOWN:
        return export_diaries(document, manifest, archive, output_dir)
    stats = dict(entries=len(document['diaries']), generated=0, skipped=0, failed=0)
    seen = set()
    for diary in document['diaries']:
        try:
            path = output_dir / diary_path(diary).with_suffix('.' + format.value)
            if path in seen:
                raise ValueError('Duplicate diary')
            seen.add(path)
            items = list(blocks(diary, manifest, archive))
            data = render_items(items, path, format)
            stats['generated' if write_exclusive(path, data) else 'skipped'] += 1
        except (OSError, ValueError, TypeError):
            stats['failed'] += 1  # Never log diary content or private paths.
    return stats


def range_filename(nickname, begin, end, format):
    import re
    from .application import validate_date_range
    validate_date_range(begin, end)
    format = ExportFormat(format)
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', nickname or '').strip().rstrip('. ')
    # Numeric/phone-like nicknames are not used as public filenames.
    if not name or re.fullmatch(r'[+\d\s()_-]+', name) or name == 'Hope 用户':
        stem = 'Hope日记'
    else:
        stem = name[:80].rstrip('. ') + '的日记'
    suffix = 'md' if format == ExportFormat.MARKDOWN else format.value
    return f'{stem}_{begin}--{end}.{suffix}'


def export_range(document, manifest, archive, output_dir, format, nickname, begin, end):
    """User-facing combined export; archive's per-entry files remain unchanged."""
    from .export_markdown import render_diary
    format = ExportFormat(format)
    archive = Path(archive).resolve()
    path = Path(output_dir).resolve() / range_filename(nickname, begin, end, format)
    stats = dict(entries=len(document['diaries']), generated=0, skipped=0, failed=0)
    if not document['diaries']:
        return stats
    try:
        if format == ExportFormat.MARKDOWN:
            data = '\n---\n\n'.join(render_diary(d, manifest, archive, path) for d in document['diaries']).encode('utf-8')
        else:
            items = []
            for diary in document['diaries']:
                if items: items.append(('separator', ''))
                items.extend(blocks(diary, manifest, archive))
            data = render_items(items, path, format)
        stats['generated' if write_exclusive(path, data) else 'skipped'] = 1
    except (OSError, ValueError, TypeError):
        stats['failed'] = 1
    return stats
