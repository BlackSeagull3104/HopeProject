"""Renderer contracts use synthetic data; no installed reader is required."""
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from hope_archive import export_documents as exports
from hope_archive.document_style import STYLES, comment_text, paragraphs
from hope_archive.image_layout import group_sizes, image_rows


def image(size):
    stream = BytesIO()
    Image.new('RGB', size, '#dce9ef').save(stream, format='PNG')
    return stream.getvalue(), size


class RenderingParityTests(unittest.TestCase):
    def test_portraits_pair_but_long_and_incompatible_do_not(self):
        for sizes, expected in [([(720,1280)]*2,[2]), ([(900,900)]*2,[2]),
                ([(720,7200)]*2,[1,1]), ([(1200,675),(720,1280)],[1,1]),
                ([(40,40)]*2,[1,1]), ([(900,900)]*3,[3]), ([(720,1280)]*4,[2,2])]:
            with self.subTest(sizes=sizes):
                self.assertEqual([count for _,count in group_sizes(sizes)],expected)

    def test_shapes_order_ratio_no_upscale_and_page_bounds(self):
        inputs = [image(size) for size in [(720,1280),(720,1280),(720,7200),(40,40),(900,900)]]
        rows = image_rows(inputs, 483.28)
        flat = [item for row in rows for item in row]
        self.assertEqual([data for data,_ in flat],[data for data,_ in inputs])
        for (_,size),(_,source) in zip(flat,inputs):
            self.assertAlmostEqual(size[0]/size[1],source[0]/source[1])
            self.assertLessEqual(size[0],source[0]*.75)
            self.assertLessEqual(size[1],360)
        for row in rows:
            self.assertLessEqual(sum(size[0] for _,size in row)+.02*483.28*(len(row)-1),483.28)

    def test_blank_lines_are_paragraphs_not_empty_word_runs(self):
        self.assertEqual(paragraphs('第一段\r\n\r\nSecond\nline\n \n最后'),['第一段','Second\nline','最后'])
        doc = Document(BytesIO(exports.render_docx([('text','第一段\n\nSecond\nline')])) )
        self.assertEqual([p.text for p in doc.paragraphs],['第一段','Second\nline'])

    def test_word_style_maps_shared_spec_without_theme_fonts(self):
        doc = Document(BytesIO(exports.render_docx([('heading','日期'),('text','中文 English')])))
        section = doc.sections[0]
        self.assertAlmostEqual(section.left_margin.pt,56)
        for name, kind in [('Normal','text'),('Heading 1','heading'),('Heading 2','title')]:
            style=doc.styles[name];spec=STYLES[kind]
            self.assertAlmostEqual(style.font.size.pt,spec.size,delta=.25)
            self.assertAlmostEqual(style.paragraph_format.line_spacing.pt,spec.leading,delta=.05)
            self.assertEqual(style.element.rPr.rFonts.get(qn('w:eastAsia')),'Microsoft YaHei')
            self.assertIsNone(style.element.rPr.rFonts.get(qn('w:asciiTheme')))
            self.assertTrue(style.paragraph_format.widow_control)

    def test_pair_and_triptych_tables_are_borderless_centered_fixed(self):
        for count in (2,3):
            doc = Document(BytesIO(exports.render_docx([('image',image((720,1280)))]*count)))
            table, = doc.tables
            self.assertEqual(len(table.columns),count*2-1)
            self.assertFalse(table.autofit)
            self.assertEqual(table.alignment,1)
            borders=table._tbl.tblPr.find(qn('w:tblBorders'))
            self.assertEqual(len(borders),6)
            self.assertTrue(all(edge.get(qn('w:val'))=='nil' for edge in borders))
            self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:cantSplit')))
            usable=(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/12700
            self.assertLessEqual(sum(column.width.pt for column in table.columns),usable+.1)

    def test_comment_author_reply_runs_and_paragraphs_preserved(self):
        comment=comment_text({'from_name':'Reader','reply_to_id':1,'text':'第一段\n\nSecond'}, {1:'Author'})
        doc=Document(BytesIO(exports.render_docx([('comment',comment)])))
        self.assertEqual([p.text for p in doc.paragraphs],['↳ Reader 回复 Author：第一段','Second'])
        self.assertEqual([r.text for r in doc.paragraphs[0].runs if r.bold],['Reader','Author'])

    def test_separator_present_in_word_and_tex(self):
        items=[('heading','First'),('text','one'),('separator',''),('heading','Last'),('text','two')]
        doc=Document(BytesIO(exports.render_docx(items)))
        self.assertEqual(len(doc.element.xpath('//w:pBdr')),1)
        with tempfile.TemporaryDirectory() as tmp:
            tex=exports.render_tex(items,Path(tmp)/'fixture.tex').decode()
            self.assertIn(r'\rule{\linewidth}',tex)
            self.assertLess(tex.index('First'),tex.index('Last'))

    def test_pdf_chinese_bold_fonts_and_determinism(self):
        items=[('heading','中文日期'),('text','第一段 English\n\n最后一段')]
        data=exports.render_pdf(items)
        self.assertEqual(data,exports.render_pdf(items))
        self.assertIn(b'/FontFile2',data)
        self.assertTrue(data.startswith(b'%PDF-'))

    def test_safe_normalized_upstream_structure(self):
        from hope_archive.normalization import normalize_diary
        diary=normalize_diary({'dairyId':1,'dairy':'fallback','noteInfo2':{'richTextInfo':[
            {'type':1,'text':'中文\n\nEnglish','fileList':None}]},
            'commentList':[{'detail':[{'commentId':1,'comments':'Synthetic comment','fromUser':{'id':1,'nickName':'Sample'}}]}]})
        items=list(exports.blocks(diary,{},Path('.')))
        doc=Document(BytesIO(exports.render_docx(items)))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('中文\nEnglish',text)
        self.assertIn('Synthetic comment',text)
        self.assertNotIn('fallback',text)


if __name__ == '__main__':
    unittest.main()
