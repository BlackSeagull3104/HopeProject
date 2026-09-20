"""Synthetic cloud/backup/archive integration; no real accounts or personal data."""
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from urllib.parse import unquote
from PIL import Image
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from hope_archive import api, archive_workflow as workflow, export_documents as docs, local_api
from hope_archive.desktop_api import DesktopService
from hope_archive.library import Library
from hope_archive.media import media_key
from hope_archive.settings import Settings


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.home=self.root/'profile';self.output=self.root/'中文归档'
        self.service=DesktopService(self.home);self.addCleanup(self.service.pool.shutdown)
        self.service.sessions['session']={'userId':'synthetic','displayName':'合成用户','expires':time.monotonic()+300}
        self.entry={'dairyId':'one','noteDate':'2024-01-01','noteType':1,'user':{'id':'synthetic'},'dairy':'中文 English synthetic diary'}

    def call(self,path,body=None,token='session'):
        return self.service.dispatch('POST',path,body or {},token)

    def configure(self): return self.call('/settings/save',{'exportRoot':str(self.output)})

    def archive(self,formats=None,entries=None):
        with patch.object(api,'fetch_all_diaries',return_value=[self.entry] if entries is None else entries):
            identity=self.call('/library/archive',{'beginDate':'2024-01-01','endDate':'2024-01-01','formats':formats or ['markdown']})['jobId']
            until=time.monotonic()+30
            while self.service.jobs[identity]['state']=='running' and time.monotonic()<until: time.sleep(.02)
            return self.service.jobs[identity]

    def test_cloud_preview_without_local_archive_and_no_disk_writes(self):
        response={'status':1,'datas':{'total':21,'list':[self.entry]}}
        with patch.object(api,'fetch_diary_page',return_value=response) as remote:
            result=self.call('/diaries/preview',{'date':'2024-01-01','diaryType':'all','page':1})
        self.assertTrue(result['hasMore']);self.assertEqual(len(result['diaries']),1)
        self.assertIsNone(remote.call_args.kwargs['raw_dir'])
        self.assertFalse(self.home.exists());self.assertFalse(self.output.exists())

    def test_cloud_filters_and_page_forwarded(self):
        from hope_archive.diary_types import entry_type,filter_value
        category=entry_type(1)
        with patch.object(api,'fetch_diary_page',return_value={'datas':{'total':21,'list':[self.entry]}}) as remote:
            self.call('/diaries/preview',{'date':'2024-01-01','diaryType':category,'page':2})
        self.assertEqual(remote.call_args.kwargs['page_num'],2)
        self.assertEqual(remote.call_args.kwargs['note_type'],filter_value(category))
        self.assertEqual(remote.call_args.args[:3],('synthetic','2024-01-01','2024-01-01'))

    def test_cloud_empty(self):
        with patch.object(api,'fetch_diary_page',return_value={'datas':{'total':0,'list':[]}}):
            result=self.call('/diaries/preview',{'date':'2024-01-01','page':1})
        self.assertEqual(result['diaries'],[]);self.assertFalse(result['hasMore'])

    def test_cloud_network_failure_is_safe(self):
        with patch.object(api,'fetch_diary_page',side_effect=api.DiaryAPIError('secret /internal/path')):
            with self.assertRaises(local_api.RequestError) as error:self.call('/diaries/preview',{'date':'2024-01-01','page':1})
        self.assertEqual(error.exception.status,502);self.assertNotIn('secret',str(error.exception))

    def test_cloud_requires_authentication(self):
        with patch.object(api,'fetch_diary_page') as remote:
            with self.assertRaises(local_api.RequestError) as error:self.call('/diaries/preview',{'date':'2024-01-01','page':1},None)
        self.assertEqual(error.exception.status,401);remote.assert_not_called()

    def test_cloud_rejects_other_owner_and_invalid_page(self):
        for body,entry in [({'date':'2024-01-01','page':0},self.entry),({'date':'2024-01-01','page':1},dict(self.entry,noteDate='2024-01-02')),({'date':'2024-01-01','page':1},dict(self.entry,user={'id':'another'}))]:
            with patch.object(api,'fetch_diary_page',return_value={'datas':{'total':1,'list':[entry]}}):
                with self.assertRaises(local_api.RequestError): self.call('/diaries/preview',body)

    def test_single_root_and_one_click_backup_archive(self):
        self.configure();job=self.archive()
        self.assertEqual(job['state'],'completed',job['stage'])
        self.assertTrue((self.output/'backup/README.txt').is_file())
        self.assertEqual(len(list((self.output/'backup').rglob('diaries.normalized.json'))),1)
        self.assertTrue(list((self.output/'backup').rglob('diaries.json')))
        self.assertEqual(len(list((self.output/'archive').glob('*.md'))),1)
        self.assertEqual(list((self.home/'work').iterdir()),[])
        self.assertFalse(list((self.output/'backup').rglob('*.md')))
        self.assertNotIn(str(self.root),json.dumps(job))

    def test_format_sets_and_single_automatic_markdown(self):
        self.configure()
        for values,expected in [(['markdown'],{'.md'}),(['pdf'],{'.md','.pdf'}),(['docx'],{'.md','.docx'}),(['pdf','docx'],{'.md','.pdf','.docx'}),(['tex'],{'.tex'})]:
            job=self.archive(values)
            self.assertEqual(job['state'],'completed',job['stage'])
            self.assertEqual({Path(n).suffix for n in job['files']},expected)
            self.assertEqual(len(job['files']),len(expected))

    def test_incremental_replaces_current_not_raw_history(self):
        self.configure();self.archive()
        self.archive(entries=[dict(self.entry,dairy='updated 合成正文')])
        paths=list((self.output/'backup').rglob('diaries.normalized.json'))
        self.assertEqual(len(paths),1)
        data=json.loads(paths[0].read_text(encoding='utf-8'))
        self.assertEqual(len(data['diaries']),1)
        self.assertEqual(data['diaries'][0]['original_text'],'updated 合成正文')
        self.assertEqual(len(list((self.output/'backup').rglob('diaries.json'))),2)
        self.assertEqual(self.call('/library/search/query',{'query':'updated'})['total'],1)

    def test_existing_media_reused_and_partial_failure_visible(self):
        self.configure()
        url='https://fixture.invalid/image.png'
        entry=dict(self.entry,noteInfo2={'richTextInfo':[{'type':'image','images':[{'url':url}]}]})
        # Patch normalization only to supply a known normalized media shape; transport remains isolated.
        from hope_archive.normalization import normalize_diaries
        document=normalize_diaries([self.entry]);document['diaries'][0]['content']=[{'kind':'image','media':[{'url':url}]}]
        account=self.service.library.account_root('synthetic');target=account/'media'/(__import__('hope_archive.media',fromlist=['local_name']).local_name(url))
        target.parent.mkdir();buffer=BytesIO();Image.new('RGB',(100,100)).save(buffer,format='PNG');target.write_bytes(buffer.getvalue())
        workflow.atomic_json(account/'media_manifest.json',{'media':{media_key(url):{'url':url,'local_path':'media/'+target.name,'status':'downloaded','sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'size':target.stat().st_size}}})
        with patch.object(workflow,'normalize_diaries',return_value=document),patch('hope_archive.media.urlopen') as network:
            result=self.archive(entries=[entry])
        self.assertEqual(result['state'],'completed',result['stage']);self.assertEqual(result['media']['skipped'],1);network.assert_not_called()
        document['diaries'][0]['content'][0]['media'].append({'url':'https://fixture.invalid/missing.png'})
        with patch.object(workflow,'normalize_diaries',return_value=document),patch('hope_archive.media.urlopen',side_effect=OSError('synthetic unavailable')):
            result=self.archive(entries=[entry])
        self.assertEqual(result['state'],'partial');self.assertIn('媒体',result['stage'])

    def test_failed_update_keeps_prior_backup_and_cleans_temp(self):
        self.configure();self.archive();path=next((self.output/'backup').rglob('diaries.normalized.json'));before=path.read_bytes()
        with patch.object(api,'fetch_all_diaries',side_effect=OSError('internal path')):
            identity=self.call('/library/archive',{'beginDate':'2024-01-01','endDate':'2024-01-01','formats':['markdown']})['jobId']
            while self.service.jobs[identity]['state']=='running':time.sleep(.01)
        self.assertEqual(path.read_bytes(),before);self.assertEqual(list((self.home/'work').iterdir()),[])
        self.assertNotIn('internal path',self.service.jobs[identity]['stage'])

    def test_search_no_legacy_path_and_fts_survives(self):
        self.configure();self.archive()
        result=self.call('/library/search/query',{'query':'synthetic'})
        self.assertEqual(result['total'],1)
        detail=self.call('/library/search/detail',{'id':result['items'][0]['id']})
        for data in (result,detail):
            self.assertNotIn('source',json.dumps(data));self.assertNotIn('normalized.json',json.dumps(data))

    def test_installation_directory_rejected(self):
        with patch.object(sys,'frozen',True,create=True),patch.object(sys,'executable',str(self.root/'install/app.exe')):
            with self.assertRaises(ValueError):Settings(self.home).save(str(self.root/'install/user-data'))

    def test_invalid_archive_formats_dates_and_account_data(self):
        self.configure()
        for formats in ([],['bad'],[{}]):
            with self.assertRaises(local_api.RequestError):
                self.call('/library/archive',{'beginDate':'2024-01-01','endDate':'2024-01-01','formats':formats})
        result=self.archive(entries=[dict(self.entry,user={'id':'other'})])
        self.assertEqual(result['state'],'failed');self.assertFalse(list((self.output/'backup').rglob('diaries.normalized.json')))

    def test_setting_change_blocked_while_archive_runs(self):
        self.configure();self.service.jobs['fixture']={'state':'running'}
        with self.assertRaises(local_api.RequestError):self.call('/settings/save',{'exportRoot':str(self.root/'other')})
        self.assertEqual(self.service.preferences.read()['exportRoot'],str(self.output))

    def test_shared_docx_tex_image_dimensions_for_all_shapes(self):
        from docx import Document
        from hope_archive.image_layout import image_rows
        sizes=[(800,450),(600,600),(400,800),(200,1800),(40,20)]
        images=[]
        for size in sizes:
            buffer=BytesIO();Image.new('RGB',size,'white').save(buffer,format='PNG');images.append((buffer.getvalue(),size))
        items=[('heading','2024-01-01'),('text','中文 English'),*[('image',im) for im in images],('comment','中文留言')]
        document=Document(BytesIO(docs.render_docx(items)))
        expected=[size for row in image_rows(images) for data,size in row]
        self.assertEqual(len(document.inline_shapes),len(sizes))
        for shape,(w,h),(ow,oh) in zip(document.inline_shapes,expected,sizes):
            self.assertAlmostEqual(shape.width.pt,w,places=3);self.assertAlmostEqual(shape.height.pt,h,places=3)
            self.assertAlmostEqual(shape.width/shape.height,ow/oh,places=4);self.assertLessEqual(shape.height.pt,300)
        tex=docs.render_tex(items,self.root/'tex/export.tex').decode()
        dimensions=[tuple(map(float,pair)) for pair in re.findall(r'width=([\d.]+)pt,height=([\d.]+)pt',tex)]
        tex_expected=[size for row in image_rows(images,453.54) for data,size in row]
        self.assertEqual(len(dimensions),5)
        for actual,expected in zip(dimensions,tex_expected):
            for a,e in zip(actual,expected):self.assertAlmostEqual(a,e,places=2)

    def test_migrated_snapshot_and_updated_canonical_search_deduplicate(self):
        library=Library(self.home);legacy=library.account_root('synthetic')/'hope-archive-old/processed/diaries.normalized.json'
        from hope_archive.normalization import normalize_diaries
        workflow.atomic_json(legacy,normalize_diaries([self.entry]))
        self.configure();self.archive(entries=[dict(self.entry,dairy='updated searchable')])
        self.assertEqual(self.call('/library/search/query',{'query':'updated'})['total'],1)
        self.assertEqual(self.call('/library/search/query',{'query':'synthetic diary'})['total'],0)
        self.assertTrue(legacy.exists())

    def test_migration_idempotent_no_overwrite_or_delete(self):
        source=self.home/'archives/account/hope-archive-old/processed/diaries.normalized.json';source.parent.mkdir(parents=True)
        source.write_text('{"diaries":[]}',encoding='utf-8')
        first=workflow.migrate_legacy(self.home,self.output/'backup')
        target=self.output/'backup'/source.relative_to(self.home/'archives')
        self.assertEqual(first['copied'],1);self.assertEqual(target.read_bytes(),source.read_bytes())
        target.write_text('{"diaries":[],"newer":true}',encoding='utf-8')
        self.assertEqual(workflow.migrate_legacy(self.home,self.output/'backup')['copied'],0)
        self.assertIn('newer',target.read_text());self.assertTrue(source.exists())

    def test_migration_conflict_preserves_both(self):
        source=self.home/'archives/account/processed/diaries.normalized.json';source.parent.mkdir(parents=True);source.write_text('old')
        target=self.output/'backup/account/processed/diaries.normalized.json';target.parent.mkdir(parents=True);target.write_text('new')
        result=workflow.migrate_legacy(self.home,self.output/'backup')
        self.assertEqual(result['conflicts'],1);self.assertEqual(target.read_text(),'new');self.assertEqual(source.read_text(),'old')

    def test_migration_failure_retains_source(self):
        source=self.home/'archives/account/processed/diaries.normalized.json';source.parent.mkdir(parents=True);source.write_text('old')
        with patch.object(shutil,'copy2',side_effect=OSError('fixture failure')):
            with self.assertRaises(OSError):workflow.migrate_legacy(self.home,self.output/'backup')
        self.assertEqual(source.read_text(),'old');self.assertFalse(list((self.output/'backup').rglob('*.migration-*')))

    def test_moved_archive_md_tex_docx_independent(self):
        backup=self.root/'backup';backup.mkdir();image=backup/'中文图片.png';Image.new('RGB',(800,500),'white').save(image)
        url='https://fixture.invalid/image';manifest={'media':{media_key(url):{'status':'downloaded','local_path':image.name}}}
        workflow.atomic_json(backup/'media_manifest.json',manifest)
        entry={'id':'one','note_date':'2024-01-01','content':[{'text':'中文 English & 100% _'}, {'kind':'image','media':[{'url':url},{'url':url}]}],'comments':[]}
        output=self.root/'archive'
        names=workflow.readable([(entry,backup)],output,workflow.formats_for(['pdf','docx','tex']),'2024-01-01','2024-01-01')
        moved=self.root/'移动后的文档';shutil.move(output,moved);shutil.rmtree(backup)
        for name in names:
            path=moved/name
            if path.suffix=='.md': refs=re.findall(r'<img src="([^"]+)"',path.read_text(encoding='utf-8'))
            elif path.suffix=='.tex': refs=re.findall(r'\\includegraphics\[[^]]+\]\{([^}]+)\}',path.read_text(encoding='utf-8'))
            else: continue
            self.assertEqual(len(refs),2)
            self.assertTrue(all((moved/unquote(ref)).is_file() for ref in refs))
        from docx import Document
        document=Document(moved/next(n for n in names if n.endswith('.docx')))
        self.assertEqual(len(document.tables),1);self.assertEqual(len(document.inline_shapes),2)
        for shape in document.inline_shapes:self.assertAlmostEqual(shape.width/shape.height,800/500,places=4)
        tex=(moved/next(n for n in names if n.endswith('.tex'))).read_text(encoding='utf-8')
        self.assertIn('fontset=fandol',tex);self.assertIn('keepaspectratio',tex);self.assertIn(r'\hspace{12pt}',tex);self.assertIn(r'100\%',tex)


if __name__=='__main__':unittest.main()
